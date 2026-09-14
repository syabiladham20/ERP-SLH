import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime

source_path = 'PPSB_PCHS_2026_Stock Report_Raw Material & Woven Bag_Harumi.xlsx'
target_path = 'PPSB_PCHS_2026_Stock_Report_Raw_Material_and_Woven_Bag_Harumi_Redesigned.xlsx'

print("Loading original workbook...")
src_wb = openpyxl.load_workbook(source_path, data_only=True)

# -------------------------------------------------------------
# 1. EXTRACT DATA FROM ORIGINAL WORKBOOK
# -------------------------------------------------------------
# A. Master Ingredients List
ingredients = []
ing_seen = set()

ws_formula = src_wb['FORMULA']
ws_sum = src_wb['SUMMARY']
ws_jan = src_wb['Jan26']

for r in range(6, ws_formula.max_row+1):
    val = ws_formula.cell(r, 1).value
    if val and isinstance(val, str):
        v = val.strip()
        if v not in ing_seen and not any(v.startswith(p) for p in ['1101', '1102', '1103', '1104', '8100', '121', '122', '123', 'PREMIX', 'TOP UP', 'WOVEN BAG', 'FORMULA', 'ADDITIVE']):
            ing_seen.add(v)
            ingredients.append(v)

for c in range(2, ws_sum.max_column+1):
    val = ws_sum.cell(3, c).value
    if val and isinstance(val, str):
        v = val.strip()
        if v not in ing_seen:
            ing_seen.add(v)
            ingredients.append(v)

for r in range(4, 50):
    val = ws_jan.cell(r, 2).value
    if val and isinstance(val, str):
        v = val.strip()
        if v not in ing_seen and v not in ['RAW MATERIAL', 'WOVEN BAG', 'MEDICINE', 'PACKAGING', 'TOTAL']:
            ing_seen.add(v)
            ingredients.append(v)

# Assign Category based on ingredient name
def get_category(name):
    u = name.upper()
    if 'MAIZE' in u or 'CORN' in u or 'WHEAT' in u or 'RICE' in u:
        return 'Grain'
    elif 'SOY' in u or 'MEAT' in u or 'PROTEIN' in u:
        return 'Protein'
    elif 'OIL' in u or 'FAT' in u:
        return 'Fat/Oil'
    elif 'BAG' in u or 'WOVEN' in u:
        return 'Packaging'
    elif 'PREMIX' in u or 'VITAMIN' in u or 'MINERAL' in u or 'CHOLINE' in u or 'TOP UP' in u:
        return 'Additive'
    else:
        return 'Micro/Raw Material'

tbl_ingredients_data = []
for i, ing in enumerate(ingredients, 1):
    code = f"ING-{i:03d}"
    unit = "KG"
    cat = get_category(ing)
    tbl_ingredients_data.append((code, ing, unit, cat))

# B. Feed Master List
feed_master_list = [
    ("1101C", "Broiler Starter 1101C", 1.0),
    ("1102C/P", "Broiler Grower 1102C/P", 1.0),
    ("1103P", "Broiler Finisher 1 1103P", 1.0),
    ("1104P", "Broiler Finisher 2 1104P", 1.0),
    ("8100", "Broiler Pre-Starter 8100", 1.0),
    ("8101", "Broiler Starter 8101", 1.0),
    ("8102", "Broiler Grower 8102", 1.0),
    ("8103", "Broiler Finisher 8103", 1.0),
    ("121", "Broiler Starter 121", 1.0),
    ("122", "Broiler Grower 122", 1.0),
    ("123", "Broiler Finisher 123", 1.0),
    ("311", "Broiler Feed 311", 1.0),
    ("312", "Broiler Feed 312", 1.0),
    ("5201", "Broiler Feed 5201", 1.0),
    ("5202", "Broiler Feed 5202", 1.0),
    ("5203", "Broiler Feed 5203", 1.0),
]

# C. Formulations Extraction
# Map feed columns in FORMULA sheet to specifications
formulation_rows = []

# Scan FORMULA sheet for headers and ingredient values
# Let's map spec headers
spec_headers = {}
current_feed_context = "1101C"

for col in range(2, ws_formula.max_column+1):
    r2 = ws_formula.cell(2, col).value
    r3 = ws_formula.cell(3, col).value
    r4 = ws_formula.cell(4, col).value
    r5 = ws_formula.cell(5, col).value

    # Check if row 2 defines a new main feed context
    if r2 and isinstance(r2, str) and ('STARTER' in r2 or 'GROWER' in r2 or 'FINISHER' in r2 or 'BROILER' in r2):
        if '1101' in r2: current_feed_context = '1101C'
        elif '1102' in r2: current_feed_context = '1102C/P'
        elif '1103' in r2: current_feed_context = '1103P'
        elif '1104' in r2: current_feed_context = '1104P'
        elif '8100' in r2 or '8101' in r2 or '8102' in r2 or '8103' in r2: current_feed_context = '8102'
        elif '121' in r2: current_feed_context = '121'
        elif '122' in r2: current_feed_context = '122'
        elif '123' in r2: current_feed_context = '123'

    spec_code = r3 or r4
    if spec_code and isinstance(spec_code, str) and (spec_code.startswith('WS') or spec_code.startswith('YH') or spec_code.startswith('31') or spec_code.startswith('52')):
        # Deduce Feed Code
        f_code = current_feed_context
        if '1101' in spec_code: f_code = '1101C'
        elif '1102' in spec_code: f_code = '1102C/P'
        elif '1103' in spec_code: f_code = '1103P'
        elif '1104' in spec_code: f_code = '1104P'
        elif '8100' in spec_code: f_code = '8100'
        elif '8101' in spec_code: f_code = '8101'
        elif '8102' in spec_code: f_code = '8102'
        elif '8103' in spec_code: f_code = '8103'
        elif '121' in spec_code: f_code = '121'
        elif '122' in spec_code: f_code = '122'
        elif '123' in spec_code: f_code = '123'
        elif '311' in spec_code: f_code = '311'
        elif '312' in spec_code: f_code = '312'
        elif '5201' in spec_code: f_code = '5201'
        elif '5202' in spec_code: f_code = '5202'
        elif '5203' in spec_code: f_code = '5203'

        # Determine formula key
        f_key = f"{f_code}-{spec_code.strip()}"

        # Check percentage column vs batch column
        col_to_check = col
        if r5 == '%' and ws_formula.cell(5, col+1).value == 'BATCH':
            col_to_check = col + 1 # Use batch column (kg per MT)

        # Iterate ingredients
        for r in range(6, ws_formula.max_row+1):
            ing_name = ws_formula.cell(r, 1).value
            inc_val = ws_formula.cell(r, col_to_check).value
            if ing_name and isinstance(ing_name, str) and inc_val is not None and isinstance(inc_val, (int, float)) and inc_val > 0:
                clean_ing = ing_name.strip()
                formulation_rows.append({
                    'Formula_Key': f_key,
                    'Feed_Code': f_code,
                    'Effective_Date': '2026-01-01',
                    'Status': 'Active',
                    'Remark': f'Standard formulation {spec_code.strip()}',
                    'Ingredient_Name': clean_ing,
                    'Inclusion_kg_per_MT': float(inc_val)
                })

print(f"Extracted {len(formulation_rows)} formulation ingredient records.")

# D. Production Orders Extraction from PRODUCE sheet
orders_rows = []
ws_prod = src_wb['PRODUCE']
# Map columns in PRODUCE sheet to Feed_Code & Spec_Code
prod_col_map = {}

# Row 2 = Feed Code, Row 3/4 = Spec Code
current_p_feed = "1101C"
for c in range(2, ws_prod.max_column+1):
    r2_v = ws_prod.cell(2, c).value
    r3_v = ws_prod.cell(3, c).value
    r4_v = ws_prod.cell(4, c).value

    if r2_v and isinstance(r2_v, str):
        if '1101' in r2_v: current_p_feed = '1101C'
        elif '1102' in r2_v: current_p_feed = '1102C/P'
        elif '1103' in r2_v: current_p_feed = '1103P'
        elif '1104' in r2_v: current_p_feed = '1104P'
        elif '8100' in r2_v or '8101' in r2_v or '8102' in r2_v or '8103' in r2_v: current_p_feed = '8102'
        elif '121' in r2_v: current_p_feed = '121'
        elif '122' in r2_v: current_p_feed = '122'
        elif '123' in r2_v: current_p_feed = '123'

    spec_v = r3_v or r4_v
    if spec_v and isinstance(spec_v, str) and any(k in spec_v for k in ['WS', 'YH', '31', '52', '81']):
        f_c = current_p_feed
        if '1101' in spec_v: f_c = '1101C'
        elif '1102' in spec_v: f_c = '1102C/P'
        elif '1103' in spec_v: f_c = '1103P'
        elif '1104' in spec_v: f_c = '1104P'
        elif '8100' in spec_v: f_c = '8100'
        elif '8101' in spec_v: f_c = '8101'
        elif '8102' in spec_v: f_c = '8102'
        elif '8103' in spec_v: f_c = '8103'
        elif '121' in spec_v: f_c = '121'
        elif '122' in spec_v: f_c = '122'
        elif '123' in spec_v: f_c = '123'

        f_k = f"{f_c}-{spec_v.strip()}"
        prod_col_map[c] = (f_c, f_k)

order_idx = 1001
for r in range(5, ws_prod.max_row+1):
    dt_val = ws_prod.cell(r, 1).value
    if dt_val and isinstance(dt_val, (datetime.datetime, datetime.date)):
        dt_str = dt_val.strftime('%Y-%m-%d')
        for col, (f_c, f_k) in prod_col_map.items():
            mt_val = ws_prod.cell(r, col).value
            if mt_val and isinstance(mt_val, (int, float)) and mt_val > 0:
                orders_rows.append({
                    'Order_ID': f"ORD-{order_idx}",
                    'Order_Date': dt_str,
                    'Feed_Code': f_c,
                    'Formula_Key': f_k,
                    'Ordered_MT': float(mt_val)
                })
                order_idx += 1

print(f"Extracted {len(orders_rows)} order records from PRODUCE sheet.")

# E. Actual Usage Extraction from VAC & MED sheet
actuals_rows = []
ws_vac = src_wb['VAC & MED']

# Scan ingredients horizontally across VAC & MED
for col in range(3, ws_vac.max_column+1, 9):
    ing_name = ws_vac.cell(2, col-1).value or ws_vac.cell(2, col).value
    if ing_name and isinstance(ing_name, str):
        clean_ing = ing_name.strip()
        for r in range(6, ws_vac.max_row+1):
            dt_val = ws_vac.cell(r, col).value
            act_val = ws_vac.cell(r, col+3).value # Actual used
            rem_val = ws_vac.cell(r, col+6).value # Remarks / Mill batch ref

            if dt_val and isinstance(dt_val, (datetime.datetime, datetime.date)) and act_val and isinstance(act_val, (int, float)) and act_val > 0:
                dt_str = dt_val.strftime('%Y-%m-%d')
                actuals_rows.append({
                    'Report_Date': dt_str,
                    'Ingredient_Name': clean_ing,
                    'Actual_Used_MT': float(act_val) / 1000.0 if float(act_val) > 100 else float(act_val), # convert kg to MT if logged in kg
                    'Mill_Batch_Ref': str(rem_val) if rem_val else 'MILL-RUN-01'
                })

print(f"Extracted {len(actuals_rows)} actual usage records from VAC & MED sheet.")

# -------------------------------------------------------------
# 2. CREATE NEW REDESIGNED WORKBOOK
# -------------------------------------------------------------
out_wb = openpyxl.Workbook()
# remove default sheet
out_wb.remove(out_wb.active)

# Styling Definitions
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Blue
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Segoe UI", size=10)
bold_font = Font(name="Segoe UI", size=10, bold=True)
title_font = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
thin_border = Border(left=Side(style='thin', color='D9D9D9'),
                     right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'),
                     bottom=Side(style='thin', color='D9D9D9'))

table_style_blue = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
table_style_green = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
table_style_orange = TableStyleInfo(name="TableStyleMedium3", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)

# SHEET 1: Master_Data
ws_master = out_wb.create_sheet(title="Master_Data")
ws_master.views.sheetView[0].showGridLines = True

# Title
ws_master['A1'] = "Master Reference Data (Ingredients & Feed Codes)"
ws_master['A1'].font = title_font

# Table 1: tbl_Ingredients (Cols A:D)
ws_master['A3'] = "Ingredient_Code"
ws_master['B3'] = "Ingredient_Name"
ws_master['C3'] = "Unit"
ws_master['D3'] = "Category"

for row_idx, data in enumerate(tbl_ingredients_data, start=4):
    for col_idx, val in enumerate(data, start=1):
        cell = ws_master.cell(row_idx, col_idx, val)
        cell.font = data_font
        cell.border = thin_border

ing_table_range = f"A3:D{len(tbl_ingredients_data)+3}"
tab_ing = Table(displayName="tbl_Ingredients", ref=ing_table_range)
tab_ing.tableStyleInfo = table_style_blue
ws_master.add_table(tab_ing)

# Table 2: tbl_Feed_Master (Cols F:H)
ws_master['F3'] = "Feed_Code"
ws_master['G3'] = "Feed_Name"
ws_master['H3'] = "Standard_Batch_Size_MT"

for row_idx, data in enumerate(feed_master_list, start=4):
    for col_idx, val in enumerate(data, start=6):
        cell = ws_master.cell(row_idx, col_idx, val)
        cell.font = data_font
        cell.border = thin_border

feed_table_range = f"F3:H{len(feed_master_list)+3}"
tab_feed = Table(displayName="tbl_Feed_Master", ref=feed_table_range)
tab_feed.tableStyleInfo = table_style_blue
ws_master.add_table(tab_feed)


# SHEET 2: Formulation_Library
ws_form = out_wb.create_sheet(title="Formulation_Library")
ws_form.views.sheetView[0].showGridLines = True

ws_form['A1'] = "Versioned Formulation Library"
ws_form['A1'].font = title_font

headers_form = ["Formula_Key", "Feed_Code", "Effective_Date", "Status", "Remark", "Ingredient_Name", "Inclusion_kg_per_MT", "Helper_Key"]
for c_idx, h in enumerate(headers_form, 1):
    cell = ws_form.cell(3, c_idx, h)
    cell.font = header_font
    cell.fill = header_fill

for r_idx, row in enumerate(formulation_rows, start=4):
    ws_form.cell(r_idx, 1, row['Formula_Key']).font = data_font
    ws_form.cell(r_idx, 2, row['Feed_Code']).font = data_font
    ws_form.cell(r_idx, 3, row['Effective_Date']).font = data_font
    ws_form.cell(r_idx, 4, row['Status']).font = data_font
    ws_form.cell(r_idx, 5, row['Remark']).font = data_font
    ws_form.cell(r_idx, 6, row['Ingredient_Name']).font = data_font

    inc_cell = ws_form.cell(r_idx, 7, row['Inclusion_kg_per_MT'])
    inc_cell.font = data_font
    inc_cell.number_format = "#,##0.00"

    # Formula for Helper_Key
    hk_cell = ws_form.cell(r_idx, 8, f'=[@Formula_Key]&"|"&[@Ingredient_Name]')
    hk_cell.font = data_font

    for c in range(1, 9):
        ws_form.cell(r_idx, c).border = thin_border

max_form_row = max(len(formulation_rows) + 3, 4)
form_table_range = f"A3:H{max_form_row}"
tab_form = Table(displayName="tbl_Formulations", ref=form_table_range)
tab_form.tableStyleInfo = table_style_green
ws_form.add_table(tab_form)


# SHEET 3: Daily_Input
ws_input = out_wb.create_sheet(title="Daily_Input")
ws_input.views.sheetView[0].showGridLines = True

ws_input['A1'] = "Daily Orders & Feedmill Actual Usage Input"
ws_input['A1'].font = title_font

# Table A: tbl_Orders (Cols A:E)
ws_input['A3'] = "Production Orders (tbl_Orders)"
ws_input['A3'].font = bold_font

headers_orders = ["Order_ID", "Order_Date", "Feed_Code", "Formula_Key", "Ordered_MT"]
for c_idx, h in enumerate(headers_orders, 1):
    cell = ws_input.cell(4, c_idx, h)
    cell.font = header_font
    cell.fill = header_fill

for r_idx, row in enumerate(orders_rows, start=5):
    ws_input.cell(r_idx, 1, row['Order_ID']).font = data_font
    ws_input.cell(r_idx, 2, row['Order_Date']).font = data_font
    ws_input.cell(r_idx, 3, row['Feed_Code']).font = data_font
    ws_input.cell(r_idx, 4, row['Formula_Key']).font = data_font

    mt_cell = ws_input.cell(r_idx, 5, row['Ordered_MT'])
    mt_cell.font = data_font
    mt_cell.number_format = "#,##0.00"

    for c in range(1, 6):
        ws_input.cell(r_idx, c).border = thin_border

max_order_row = max(len(orders_rows) + 4, 5)
orders_table_range = f"A4:E{max_order_row}"
tab_orders = Table(displayName="tbl_Orders", ref=orders_table_range)
tab_orders.tableStyleInfo = table_style_blue
ws_input.add_table(tab_orders)

# Add Data Validation dropdown for Formula_Key in tbl_Orders
dv_formula = DataValidation(type="list", formula1="=Formulation_Library!$A$4:$A$" + str(max_form_row), allow_blank=True)
ws_input.add_data_validation(dv_formula)
dv_formula.add(f"D5:D{max_order_row}")


# Table B: tbl_Feedmill_Actuals (Cols G:J)
ws_input['G3'] = "Feedmill Actual Usage (tbl_Feedmill_Actuals)"
ws_input['G3'].font = bold_font

headers_actuals = ["Report_Date", "Ingredient_Name", "Actual_Used_MT", "Mill_Batch_Ref"]
for c_idx, h in enumerate(headers_actuals, start=7):
    cell = ws_input.cell(4, c_idx, h)
    cell.font = header_font
    cell.fill = header_fill

for r_idx, row in enumerate(actuals_rows, start=5):
    ws_input.cell(r_idx, 7, row['Report_Date']).font = data_font
    ws_input.cell(r_idx, 8, row['Ingredient_Name']).font = data_font

    act_cell = ws_input.cell(r_idx, 9, row['Actual_Used_MT'])
    act_cell.font = data_font
    act_cell.number_format = "#,##0.00"

    ws_input.cell(r_idx, 10, row['Mill_Batch_Ref']).font = data_font

    for c in range(7, 11):
        ws_input.cell(r_idx, c).border = thin_border

max_actual_row = max(len(actuals_rows) + 4, 5)
actuals_table_range = f"G4:J{max_actual_row}"
tab_actuals = Table(displayName="tbl_Feedmill_Actuals", ref=actuals_table_range)
tab_actuals.tableStyleInfo = table_style_orange
ws_input.add_table(tab_actuals)


# SHEET 4: Usage_Reconciliation
ws_recon = out_wb.create_sheet(title="Usage_Reconciliation")
ws_recon.views.sheetView[0].showGridLines = True

ws_recon['A1'] = "Daily & Monthly Raw Material Usage Reconciliation"
ws_recon['A1'].font = title_font

# Selector Controls
ws_recon['A3'] = "Target Date:"
ws_recon['A3'].font = bold_font
ws_recon['B3'] = "2026-01-02"
ws_recon['B3'].font = bold_font
ws_recon['B3'].alignment = Alignment(horizontal="center")
ws_recon['B3'].fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Soft Yellow Highlight

ws_recon['A4'] = "Month Start Date:"
ws_recon['A4'].font = bold_font
ws_recon['B4'] = "=EOMONTH(B3, -1) + 1"
ws_recon['B4'].font = bold_font
ws_recon['B4'].alignment = Alignment(horizontal="center")

headers_recon = [
    "Ingredient_Name",
    "Day Expected (MT)",
    "Day Actual (MT)",
    "Day Var (MT)",
    "Day Var %",
    "MTD Expected (MT)",
    "MTD Actual (MT)",
    "MTD Var (MT)",
    "MTD Var %"
]

for c_idx, h in enumerate(headers_recon, start=1):
    cell = ws_recon.cell(6, c_idx, h)
    cell.font = header_font
    cell.fill = header_fill

for r_idx, ing_data in enumerate(tbl_ingredients_data, start=7):
    ing_name = ing_data[1]
    ws_recon.cell(r_idx, 1, ing_name).font = data_font

    # Day Expected (MT)
    # SUM(FILTER(tbl_Orders[Ordered_MT], tbl_Orders[Order_Date] = $B$3, 0) * XLOOKUP(FILTER(tbl_Orders[Formula_Key], tbl_Orders[Order_Date] = $B$3, "") & "|" & [@Ingredient_Name], tbl_Formulations[Helper_Key], tbl_Formulations[Inclusion_kg_per_MT], 0) / 1000)
    day_exp_f = f'=SUM(FILTER(tbl_Orders[Ordered_MT], tbl_Orders[Order_Date] = $B$3, 0) * XLOOKUP(FILTER(tbl_Orders[Formula_Key], tbl_Orders[Order_Date] = $B$3, "") & "|" & [@Ingredient_Name], tbl_Formulations[Helper_Key], tbl_Formulations[Inclusion_kg_per_MT], 0) / 1000)'
    cell_de = ws_recon.cell(r_idx, 2, day_exp_f)
    cell_de.font = data_font
    cell_de.number_format = "#,##0.00"

    # Day Actual (MT)
    day_act_f = f'=SUMIFS(tbl_Feedmill_Actuals[Actual_Used_MT], tbl_Feedmill_Actuals[Report_Date], $B$3, tbl_Feedmill_Actuals[Ingredient_Name], [@Ingredient_Name])'
    cell_da = ws_recon.cell(r_idx, 3, day_act_f)
    cell_da.font = data_font
    cell_da.number_format = "#,##0.00"

    # Day Var (MT)
    cell_dv = ws_recon.cell(r_idx, 4, f'=[@[Day Actual (MT)]]-[@[Day Expected (MT)]]')
    cell_dv.font = data_font
    cell_dv.number_format = "#,##0.00"

    # Day Var %
    cell_dvp = ws_recon.cell(r_idx, 5, f'=IFERROR([@[Day Var (MT)]]/[@[Day Expected (MT)]], 0)')
    cell_dvp.font = data_font
    cell_dvp.number_format = "0.00%"

    # MTD Expected (MT)
    mtd_exp_f = f'=LET(in_range, (tbl_Orders[Order_Date] >= $B$4) * (tbl_Orders[Order_Date] <= $B$3), filtered_mt, FILTER(tbl_Orders[Ordered_MT], in_range, 0), filtered_keys, FILTER(tbl_Orders[Formula_Key], in_range, ""), SUM(filtered_mt * XLOOKUP(filtered_keys & "|" & [@Ingredient_Name], tbl_Formulations[Helper_Key], tbl_Formulations[Inclusion_kg_per_MT], 0) / 1000))'
    cell_me = ws_recon.cell(r_idx, 6, mtd_exp_f)
    cell_me.font = data_font
    cell_me.number_format = "#,##0.00"

    # MTD Actual (MT)
    mtd_act_f = f'=SUMIFS(tbl_Feedmill_Actuals[Actual_Used_MT], tbl_Feedmill_Actuals[Report_Date], ">=" & $B$4, tbl_Feedmill_Actuals[Report_Date], "<=" & $B$3, tbl_Feedmill_Actuals[Ingredient_Name], [@Ingredient_Name])'
    cell_ma = ws_recon.cell(r_idx, 7, mtd_act_f)
    cell_ma.font = data_font
    cell_ma.number_format = "#,##0.00"

    # MTD Var (MT)
    cell_mv = ws_recon.cell(r_idx, 8, f'=[@[MTD Actual (MT)]]-[@[MTD Expected (MT)]]')
    cell_mv.font = data_font
    cell_mv.number_format = "#,##0.00"

    # MTD Var %
    cell_mvp = ws_recon.cell(r_idx, 9, f'=IFERROR([@[MTD Var (MT)]]/[@[MTD Expected (MT)]], 0)')
    cell_mvp.font = data_font
    cell_mvp.number_format = "0.00%"

    for c in range(1, 10):
        ws_recon.cell(r_idx, c).border = thin_border

max_recon_row = len(tbl_ingredients_data) + 6
recon_table_range = f"A6:I{max_recon_row}"
tab_recon = Table(displayName="tbl_Usage_Reconciliation", ref=recon_table_range)
tab_recon.tableStyleInfo = table_style_blue
ws_recon.add_table(tab_recon)


# -------------------------------------------------------------
# 3. COPY HISTORICAL SHEETS FROM ORIGINAL WORKBOOK
# -------------------------------------------------------------
print("Copying historical sheets from original workbook...")
sheets_to_copy = ['FORMULA', 'PRODUCE', 'VAC & MED', 'SUMMARY', 'Jan26', 'Feb26', 'Mar26', 'Apr26', 'May26', 'Jun26', 'Jul26', 'Aug26', 'Sept26', 'Oct26', 'Nov26', 'Dec26']

for sname in sheets_to_copy:
    if sname in src_wb.sheetnames:
        src_sheet = src_wb[sname]
        dst_sheet = out_wb.create_sheet(title=sname)
        dst_sheet.views.sheetView[0].showGridLines = True

        for r in range(1, src_sheet.max_row+1):
            for c in range(1, src_sheet.max_column+1):
                cell_val = src_sheet.cell(r, c).value
                if cell_val is not None:
                    dst_cell = dst_sheet.cell(r, c, cell_val)
                    dst_cell.font = data_font

# Adjust column widths across all sheets
print("Adjusting column widths...")
for sheet in out_wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col[:100]: # check top 100 rows
            if cell.value:
                val_str = str(cell.value)
                if not val_str.startswith('='):
                    max_len = max(max_len, len(val_str))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

print(f"Saving redesigned workbook to {target_path}...")
out_wb.save(target_path)
print("Workbook saved successfully!")
